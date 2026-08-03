"""
Export the WHOLE benchmark into ONE table (one CSV + one single-sheet Excel),
so tasks, summary metrics, per-document results, and every ground-truth label
aligned with its corresponding prediction are all visible in one place.

Every row shares one schema (the `section` column says what kind of row it is):

  section   | what the row is                | key columns used
  ----------+--------------------------------+--------------------------------
  info      | run config                     | key, ground_truth(=value)
  task      | a project task                 | key(id), ground_truth(subject), details(group), verdict(status)
  summary   | overall / per-class metric     | key(name), precision..fn
  document  | per-document metric            | key(doc_id), link, precision..fn
  detail    | gold label ↔ prediction        | key(doc_id), link, class, kind,
            |   (the aligned comparison)     |   ground_truth, matched, prediction, verdict, details(attrs)

Usage:
  cd query_engine
  python -m extract_eval.regression.export_combined report.json benchmark_all
  #   → benchmark_all.csv  and  benchmark_all.xlsx
  # or via the runner:  evaluate_regression ... --combined benchmark_all
"""
from __future__ import annotations

import csv
import json
import sys

from extract_eval.regression.scorer import SCORED_CLASSES, INFORMATIONAL_CLASSES, matches

COLUMNS = [
    "section", "key", "class", "link", "kind", "ground_truth", "matched",
    "prediction", "verdict", "precision", "recall", "f1", "tp", "fp", "fn", "details",
]

# Default project tasks (this session). Override by passing tasks= to build_rows.
DEFAULT_TASKS = [
    (1, "Update build_db.py schema", "Attack/Impact ontology", "completed"),
    (2, "Update build_rdf.py", "Attack/Impact ontology", "completed"),
    (3, "Make Attack/Impact extractable in ontology.py", "Attack/Impact ontology", "completed"),
    (4, "Wire Attack/Impact into persist.py", "Attack/Impact ontology", "completed"),
    (5, "Update app.py API + prompts", "Attack/Impact ontology", "completed"),
    (6, "Migrate DB, regen RDF, cleanup, test", "Attack/Impact ontology", "completed"),
    (7, "Refactor pipeline for markdown input", "Regression benchmark", "completed"),
    (8, "Make base LLM swappable via env", "Regression benchmark", "completed"),
    (9, "Build crawl fixtures from test_links.csv", "Regression benchmark", "completed"),
    (10, "Oracle-label ground truth gold.yaml", "Regression benchmark", "completed"),
    (11, "Build scorer + evaluation runner", "Regression benchmark", "completed"),
    (12, "Run benchmark, report P/R/F1, document", "Regression benchmark", "completed"),
]


def _round(x):
    return "" if x is None else round(x, 4)


def _reconstruct_item(class_name: str, gold_row: dict) -> dict:
    """Rebuild a matchable gold item from a stored gold detail row."""
    label = gold_row["label"]
    if class_name == "Vulnerability":
        if label.startswith("title_any:"):
            return {"title_any": label[len("title_any:"):].split("|")}
        return {"id": label}
    return {"name": label, "aliases": gold_row.get("aliases", [])}


def _pval(p: dict) -> str:
    """Display value for a predicted entity (falls back to id for CWE types)."""
    a = p.get("attrs", {})
    return p.get("value") or a.get("name") or a.get("vuln_id") or a.get("id") or a.get("title") or ""


def _row(**kw) -> dict:
    r = {c: "" for c in COLUMNS}
    r.update(kw)
    return r


def build_rows(report: dict, tasks=DEFAULT_TASKS) -> list[dict]:
    rows: list[dict] = []
    cfg = report.get("config", {})

    # ── run info ──
    for k in ("generated_at", "model", "source", "n_docs"):
        rows.append(_row(section="info", key=k, ground_truth=cfg.get(k, "")))

    # ── tasks ──
    for tid, subject, group, status in (tasks or []):
        rows.append(_row(section="task", key=tid, ground_truth=subject,
                         details=group, verdict=status))

    # ── summary metrics ──
    o = report["overall"]
    rows.append(_row(section="summary", key="OVERALL (headline)",
                     precision=_round(o["precision"]), recall=_round(o["recall"]),
                     f1=_round(o["f1"]), tp=o["tp"], fp=o["fp"], fn=o["fn"]))
    for c, m in report["per_class"].items():
        rows.append(_row(section="summary", key=c, class_=c,
                         precision=_round(m["precision"]), recall=_round(m["recall"]),
                         f1=_round(m["f1"]), tp=m["tp"], fp=m["fp"], fn=m["fn"],
                         details=("informational — excluded from OVERALL"
                                  if c in INFORMATIONAL_CLASSES else "")))

    # ── per-document metrics ──
    for doc_id, d in report["per_doc"].items():
        meta = d.get("meta", {})
        m = d.get("metrics") or {}
        rows.append(_row(section="document", key=doc_id, link=meta.get("url", ""),
                         precision=_round(m.get("precision")), recall=_round(m.get("recall")),
                         f1=_round(m.get("f1")), tp=m.get("tp"), fp=m.get("fp"), fn=m.get("fn"),
                         ground_truth=meta.get("title", "")))

    # ── detail: every gold label aligned with its matching prediction ──
    for doc_id, d in report["per_doc"].items():
        meta = d.get("meta", {})
        url = meta.get("url", "")
        details = meta.get("details") or {}
        for c in SCORED_CLASSES:
            preds = details.get(c, {}).get("predicted", [])
            gold = details.get(c, {}).get("gold", [])
            used = set()
            for g in gold:
                item = _reconstruct_item(c, g)
                hits = [p for p in preds if matches(c, p.get("attrs", {}), item)]
                if hits:
                    for p in hits:
                        used.add(_pval(p))
                        rows.append(_row(section="detail", key=doc_id, class_=c, link=url,
                                         kind=g["kind"], ground_truth=g["label"], matched="yes",
                                         prediction=_pval(p), verdict=p["verdict"],
                                         details=json.dumps(p.get("attrs", {}), ensure_ascii=False)))
                else:
                    rows.append(_row(section="detail", key=doc_id, class_=c, link=url,
                                     kind=g["kind"], ground_truth=g["label"], matched="no",
                                     verdict=("MISS" if g["kind"] == "required" else "")))
            # predictions that matched no gold item = false positives
            for p in preds:
                if _pval(p) in used:
                    continue
                if p["verdict"] == "FP":
                    rows.append(_row(section="detail", key=doc_id, class_=c, link=url,
                                     kind="", ground_truth="(no matching gold)", matched="",
                                     prediction=_pval(p), verdict="FP",
                                     details=json.dumps(p.get("attrs", {}), ensure_ascii=False)))
    return rows


def _fix_class_key(rows):
    # allow class_=... kwarg → 'class' column
    for r in rows:
        if "class_" in r:
            r["class"] = r.pop("class_")
    return rows


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    section_fill = {
        "info": "DDEBF7", "task": "E2EFDA", "summary": "FCE4D6",
        "document": "FFF2CC", "detail": "FFFFFF",
    }
    verdict_fill = {"TP": "C6EFCE", "acceptable": "FFEB9C", "FP": "FFC7CE", "MISS": "FFC7CE"}
    wb = Workbook()
    ws = wb.active
    ws.title = "benchmark"
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    for j, col in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = head_fill; c.font = head_font
    ws.freeze_panes = "A2"
    link_i, verdict_i, section_i = COLUMNS.index("link") + 1, COLUMNS.index("verdict") + 1, COLUMNS.index("section") + 1
    for i, r in enumerate(rows, start=2):
        for j, col in enumerate(COLUMNS, 1):
            ws.cell(row=i, column=j, value=r.get(col, ""))
        sc = ws.cell(row=i, column=section_i)
        sc.fill = PatternFill("solid", fgColor=section_fill.get(r["section"], "FFFFFF"))
        sc.font = Font(bold=True)
        url = r.get("link")
        if url:
            lc = ws.cell(row=i, column=link_i)
            lc.hyperlink = url; lc.font = Font(color="0563C1", underline="single")
        vf = verdict_fill.get(r.get("verdict"))
        if vf:
            ws.cell(row=i, column=verdict_i).fill = PatternFill("solid", fgColor=vf)
    widths = [10, 40, 16, 46, 12, 34, 9, 30, 11, 9, 9, 8, 5, 5, 5, 60]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    wb.save(path)


def export(report: dict, out_base: str, tasks=DEFAULT_TASKS) -> list[str]:
    rows = _fix_class_key(build_rows(report, tasks))
    csv_path, xlsx_path = out_base + ".csv", out_base + ".xlsx"
    write_csv(rows, csv_path)
    write_xlsx(rows, xlsx_path)
    return [csv_path, xlsx_path]


def main() -> None:
    if len(sys.argv) < 3:
        print("usage: python -m extract_eval.regression.export_combined <report.json> <out_base>")
        raise SystemExit(2)
    with open(sys.argv[1], encoding="utf-8") as f:
        report = json.load(f)
    for p in export(report, sys.argv[2]):
        print("wrote", p)


if __name__ == "__main__":
    main()
